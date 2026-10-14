"""
文档加载器模块 —— 支持 pdf / word / excel / ppt / csv / html / json / txt / markdown 等格式的加载与切分。

职责：
1. 加载文档：根据文件后缀选择对应的 Loader 解析，支持单个文件与整目录（含子目录）批量加载。
2. 注入元数据：为每个片段补充「文档类型 file_type、文件名 file_name、文件路径 source」，
   便于后续溯源展示与按来源删除。
3. 分割文档：使用 RecursiveCharacterTextSplitter 切成大小合适的片段，供向量化入库。

设计要点：
- 所有入口都会做边界校验：路径必须存在、必须是对应的文件/目录、后缀必须在白名单内。
- 单个文件失败只记录日志并跳过，不会中断整目录的批量加载。
- 全链路关键节点均有日志（校验、解析、切分、汇总）。
"""
import json
import logging
from collections.abc import Callable
from pathlib import Path

from langchain_community.document_loaders import (
    BSHTMLLoader,
    CSVLoader,
    Docx2txtLoader,
    PyMuPDFLoader,
    TextLoader,
    UnstructuredExcelLoader,
    UnstructuredMarkdownLoader,
)
from langchain_community.document_loaders.base import BaseLoader
from pptx import Presentation
from pptx.table import Table
from pptx.text.text import TextFrame
from langchain_core.documents import Document
from langchain_text_splitters import MarkdownHeaderTextSplitter, RecursiveCharacterTextSplitter

from config.settings import settings

logger = logging.getLogger(__name__)


class DocumentLoader:
    """统一文档加载器：负责「文件 → 文本片段列表」的全过程。"""

    # 全部支持的后缀白名单（扩展名统一小写，含前缀点）；路径校验与目录遍历都以此为准
    SUPPORTED_EXTENSIONS: frozenset[str] = frozenset({
        ".pdf", ".doc", ".docx", ".txt", ".md", ".xlsx", ".xls", ".pptx",
        ".csv", ".json", ".html", ".htm",
    })

    # 可直接以「路径」为唯一必填入参实例化的 Loader；
    # txt/md/csv/html 需要编码回退、json 需要结构化解析，统一走 _load_custom 分支
    LOADER_MAP: dict[str, Callable[[str], BaseLoader]] = {
        ".pdf": PyMuPDFLoader,
        ".docx": Docx2txtLoader,
        ".doc": Docx2txtLoader,
        ".xlsx": UnstructuredExcelLoader,
        ".xls": UnstructuredExcelLoader,
    }

    # 需要自定义解析逻辑的后缀：编码回退 / 结构化 schema / 绕开 unstructured 的分块依赖
    CUSTOM_EXTENSIONS: frozenset[str] = frozenset({
        ".txt", ".md", ".csv", ".json", ".html", ".htm", ".pptx",
    })

    # 旧版二进制 Office 格式：解析库支持有限，命中时给出明确告警，便于排查
    LEGACY_BINARY_FORMATS: set[str] = {".doc", ".xls"}

    # 纯文本类文件解码时按顺序尝试的编码（中文文档常见 gbk 系列）
    TEXT_ENCODINGS: tuple[str, ...] = ("utf-8", "utf-8-sig", "gbk", "gb18030")

    # 遍历目录时忽略的文件名前缀：隐藏文件与 Office 临时文件
    IGNORED_NAME_PREFIXES: tuple[str, ...] = (".", "~$")

    def __init__(self, chunk_size: int | None = None, chunk_overlap: int | None = None):
        """
        :param chunk_size: 每个片段的字符数，缺省读取 settings.CHUNK_SIZE
        :param chunk_overlap: 相邻片段的重叠字符数，缺省读取 settings.CHUNK_OVERLAP
        """
        self.chunk_size: int = chunk_size or settings.CHUNK_SIZE
        self.chunk_overlap: int = chunk_overlap or settings.CHUNK_OVERLAP
        # 切分模式（P3-7a）：fixed = 现状递归切分；title = 标题感知 + Enrichment。
        # 只在 md（含 B 阶段 MinerU 产出的 markdown）上生效；其余格式一律走 fixed 路径。
        self.chunk_mode: str = settings.CHUNK_MODE
        self.text_splitter: RecursiveCharacterTextSplitter = RecursiveCharacterTextSplitter(
            chunk_size=self.chunk_size,
            chunk_overlap=self.chunk_overlap,
            # 中文优先的分隔符：先按段落/换行，再按中文句末标点，最后按英文标点与空格
            separators=["\n\n", "\n", "。", "！", "？", ".", "!", "?", " ", ""],
            length_function=len,
        )
        logger.info(
            "DocumentLoader 初始化完成 | chunk_size=%s chunk_overlap=%s 支持类型=%s",
            self.chunk_size,
            self.chunk_overlap,
            "/".join(sorted(ext.lstrip(".") for ext in self.SUPPORTED_EXTENSIONS)),
        )

    # ------------------------------------------------------------------ #
    # 公共能力
    # ------------------------------------------------------------------ #
    @classmethod
    def is_supported(cls, file_path: str | Path) -> bool:
        """判断文件后缀是否在支持白名单内（仅看后缀，不校验文件是否存在）。"""
        return Path(file_path).suffix.lower() in cls.SUPPORTED_EXTENSIONS

    @staticmethod
    def find_duplicate_names(documents: list[Document]) -> dict[str, list[str]]:
        """
        找出「文件名相同、但来源路径不同」的文档。

        这类文件内容可能完全不同（如 A 目录与 B 目录各有一个「说明.txt」），
        加载层不会丢弃任何一份，但调用方需要感知，以免按文件名删除 / 展示时张冠李戴。

        :param documents: 片段列表（load_file / load_directory 的返回值）
        :return: {文件名: [来源路径, ...]}，只包含同名不同路径的条目；无冲突时为空字典
        """
        name_to_sources: dict[str, set[str]] = {}
        for doc in documents:
            name = str(doc.metadata.get("file_name", ""))
            source = str(doc.metadata.get("source", ""))
            if name:
                name_to_sources.setdefault(name, set()).add(source)
        return {
            name: sorted(sources)
            for name, sources in name_to_sources.items()
            if len(sources) > 1
        }

    # ------------------------------------------------------------------ #
    # 单个文件加载
    # ------------------------------------------------------------------ #
    def load_file(self, file_path: str | Path, metadata: dict[str, object] | None = None) -> list[Document]:
        """
        加载并切分单个文件。

        :param file_path: 文件路径
        :param metadata: 可选的额外元数据，会与内置元数据合并（同名字段以调用方传入为准）
        :return: 切分后的片段列表；任何校验不通过或解析失败均返回空列表
        """
        path = Path(file_path)

        # 1. 路径校验：必须存在，且必须是文件（目录、死链等一律拒绝）
        if not path.exists():
            logger.error("加载失败：文件不存在 -> %s", path)
            return []
        if not path.is_file():
            logger.error("加载失败：路径不是文件 -> %s", path)
            return []

        # 2. 后缀校验：必须有后缀，且必须在白名单内
        ext = path.suffix.lower()
        if not ext:
            logger.error("加载失败：文件缺少扩展名 -> %s", path)
            return []
        if ext not in self.SUPPORTED_EXTENSIONS:
            logger.error(
                "加载失败：不支持的文件类型 '%s' -> %s；当前支持：%s",
                ext,
                path,
                "/".join(sorted(e.lstrip(".") for e in self.SUPPORTED_EXTENSIONS)),
            )
            return []

        # 3. 组装内置元数据：文档类型、文件名、文件路径
        base_metadata: dict[str, object] = {
            "source": str(path),          # 文件路径（用于溯源与按来源删除）
            "file_name": path.name,       # 文件名
            "file_type": ext.lstrip("."), # 文档类型（不含点）
        }
        if metadata:
            logger.debug("合并调用方元数据 %s -> %s", metadata, path)
            base_metadata.update(metadata)

        # 4. 解析文档
        # title 模式（P3-7a）只对 md 开放闭环：直接读原始 markdown 文本按标题树切，
        # 不经过 Loader（UnstructuredMarkdownLoader 会把 `#` 标记吃掉，标题树无从谈起）。
        # 失败回退 fixed 路径 —— 切分模式是增强，不该让一篇坏文件整批失败。
        if self.chunk_mode == "title" and ext == ".md":
            try:
                return self._load_md_title_mode(path, base_metadata)
            except Exception as e:
                logger.error("标题模式切分失败，回退固定切分：%s | 错误：%s", path, e, exc_info=True)

        try:
            documents = self._load_documents(path, ext)
        except Exception as e:  # 解析失败只影响当前文件，向上返回空列表
            logger.error("解析文件失败：%s | 错误：%s", path, e, exc_info=True)
            return []

        if not documents:
            logger.warning("文件解析结果为空：%s", path)
            return []
        logger.info("解析完成：%s | 原始文档段数=%d", path, len(documents))

        # 5. 注入元数据（Loader 自带的元数据如 page 会保留，同名以自定义为准）
        for doc in documents:
            doc.metadata.update(base_metadata)

        # 6. 切分文档
        try:
            split_docs = self.text_splitter.split_documents(documents)
        except Exception as e:
            logger.error("切分文档失败：%s | 错误：%s", path, e, exc_info=True)
            return []

        logger.info(
            "加载并切分完成：%s | 原始段数=%d 片段数=%d",
            path,
            len(documents),
            len(split_docs),
        )
        return split_docs

    # ------------------------------------------------------------------ #
    # title 模式（P3-7a）：标题感知切分 + Enrichment
    # ------------------------------------------------------------------ #
    # 参与标题树切分的层级：h1~h4 再深的小节通常已是段落粒度，切到那一层
    # 只会把连贯段落拆散。5 级以后的小节留给「超长二次切分」去处理。
    TITLE_HEADER_LEVELS: int = 4
    # 过短小节阈值（相对 chunk_size）：低于它的 section 向相邻合并。
    # 0.3 是「一段话 vs 一个 chunk」的经验分界 —— 150 字的「3.2 电压要求」
    # 单独成片，embedding 里全是标题、正文信息量不足，不如并给邻居。
    SHORT_SECTION_RATIO: float = 0.3

    def _load_md_title_mode(self, path: Path, base_metadata: dict[str, object]) -> list[Document]:
        """
        markdown 标题感知切分：标题树分节 → 短节合并 → 长节二次切 → Enrichment 入库。

        每个最终片段：
            page_content        = enriched（header_path + 原文）—— 进 embedding 与 Chroma 正文；
            metadata["content"] = 原文（引用反查 / 前端展示用它，不让用户看到拼接串）；
            metadata["header_path"] = 「一级 > 二级 > 三级」章节路径。

        为什么 enriched 放 page_content 而不是「预计算向量再存原文」：
        langchain 的 Chroma.add_documents 只会嵌 page_content，注入预计算向量
        要绕过它直连 collection.add，等于把 P0-3 收敛的写入路径再抄一份 —— 不值。
        代价是读取侧（vector_store / rag_chain）要认 metadata["content"]，
        这个约定由 VectorStoreManager._display_content() 单点承担。
        """
        md_text = self._read_text(path)
        header_splitter = MarkdownHeaderTextSplitter(
            headers_to_split_on=[
                ("#", "header_1"), ("##", "header_2"),
                ("###", "header_3"), ("####", "header_4"),
            ],
            strip_headers=False,   # 标题行留在正文里：原文完整性 + 关键词可见
        )
        sections = header_splitter.split_text(md_text)
        logger.info("标题模式分节完成：%s | 小节数=%d", path, len(sections))

        # ---- 1. 短节向相邻合并（先向后找归宿，不落单）----
        # 合并块的 header_path 取两节标题的**最深公共祖先**：
        # 「一、发放规则」并进「二、工资条查询」的块时，谁都吞不掉谁，
        # 诚实标注是共同父级「工资查询指引」—— 否则 enrichment 前缀会撒谎，
        # embedding 学到「发放规则的内容叫常见问题」这种错位语义。
        short_threshold = int(self.chunk_size * self.SHORT_SECTION_RATIO)
        blocks: list[tuple[str, str]] = []      # (header_path, section_text)
        pending_short: tuple[str, str] | None = None
        for sec in sections:
            text = sec.page_content.strip()
            if not text:
                continue
            header_path = self._header_path_of(sec.metadata)
            if len(text) < short_threshold:
                if blocks:
                    # 并入上一块：header 取两节的公共祖先（内容上是延续+新话题）
                    prev_hp, prev_text = blocks[-1]
                    blocks[-1] = (
                        self._common_ancestor(prev_hp, header_path),
                        f"{prev_text}\n\n{text}",
                    )
                elif pending_short is not None:
                    # 开头连续多个短节：**累加**而不是覆盖 ——
                    # 覆盖会把前面的短节静默丢掉（v3.1.0-p3.7a 自测抓到的真 bug）
                    php, pst = pending_short
                    pending_short = (
                        self._common_ancestor(php, header_path),
                        f"{pst}\n\n{text}",
                    )
                else:
                    pending_short = (header_path, text)
                continue
            if pending_short is not None:
                hp, st = pending_short
                text = f"{st}\n\n{text}"
                header_path = self._common_ancestor(hp, header_path)
                pending_short = None
            blocks.append((header_path, text))
        if pending_short is not None:
            blocks.append(pending_short)        # 整篇都短：单成一片，总比丢掉强

        # ---- 2. 长节二次切（表格原子）+ Enrichment ----
        documents: list[Document] = []
        for header_path, text in blocks:
            for piece in self._split_block_protect_tables(text):
                enriched = f"{header_path}\n{piece}" if header_path else piece
                documents.append(Document(
                    page_content=enriched,
                    metadata={
                        **base_metadata,
                        "header_path": header_path,
                        "content": piece,       # 原文，读取侧展示/引用反查用
                    },
                ))

        logger.info(
            "标题模式切分完成：%s | 小节=%d 片段=%d（短节阈值=%d）",
            path, len(sections), len(documents), short_threshold,
        )
        return documents

    @classmethod
    def _header_path_of(cls, meta: dict) -> str:
        """把 MarkdownHeaderTextSplitter 给的 header_1..4 元数据拼成「A > B > C」。"""
        parts = [
            str(meta[key]).strip()
            for key in ("header_1", "header_2", "header_3", "header_4")[: cls.TITLE_HEADER_LEVELS]
            if meta.get(key)
        ]
        return " > ".join(parts)

    @staticmethod
    def _common_ancestor(hp1: str, hp2: str) -> str:
        """两条 header_path 的最深公共祖先（前缀一致的最长一条）。"""
        if not hp1:
            return hp2
        if not hp2:
            return hp1
        a, b = hp1.split(" > "), hp2.split(" > ")
        parts: list[str] = []
        for x, y in zip(a, b):
            if x != y:
                break
            parts.append(x)
        return " > ".join(parts)

    def _split_block_protect_tables(self, text: str) -> list[str]:
        """
        把一个（可能超长的）小节切成 ≤ chunk_size 的片段，**表格行绝不拦腰断**。

        表格（连续以 | 开头的行）处理分两档：
            · 整表 ≤ chunk_size → 整块独成一个片段（合并单元格/跨页拼接的完整网格
              只有整存才答得了「XX 行 YY 列」）；
            · 整表 > chunk_size → 按**行组**切：每个分片重复表头行 + 分隔行，
              行本身绝不切断。直接整块塞给 embedding 会撞嵌入模型输入上限
              （bge-large-zh-v1.5 约 512 token，实测 773 字的表直接 400）——
              「原子」的本意是行级完整 + 表头随行，不是无限长的单块。
        普通段落超过 chunk_size 时才用既有 RecursiveCharacterTextSplitter 二次切。
        """
        # 先把文本切成 (是否表格行, 行列表) 的连续段
        segments: list[tuple[bool, list[str]]] = []
        cur_is_table = False
        cur: list[str] = []
        for line in text.split("\n"):
            is_table_row = line.lstrip().startswith("|")
            if is_table_row != cur_is_table and cur:
                segments.append((cur_is_table, cur))
                cur = []
            cur_is_table = is_table_row
            cur.append(line)
        if cur:
            segments.append((cur_is_table, cur))

        pieces: list[str] = []
        buf: list[str] = []
        buf_len = 0

        def flush() -> None:
            nonlocal buf, buf_len
            if buf:
                joined = "\n\n".join(buf).strip()
                if joined:
                    pieces.append(joined)
            buf, buf_len = [], 0

        for is_table, lines in segments:
            seg_text = "\n".join(lines).strip()
            if not seg_text:
                continue
            if is_table:
                flush()
                if len(seg_text) <= self.chunk_size:
                    pieces.append(seg_text)     # 整表原子（≤ chunk_size）
                else:
                    # 超长表按行组切：头两行（表头 + 分隔行）在每个分片重复
                    header = lines[:2]
                    header_len = len("\n".join(header)) + 1
                    group: list[str] = []
                    group_len = header_len
                    for row in lines[2:]:
                        row_len = len(row) + 1
                        if group and group_len + row_len > self.chunk_size:
                            pieces.append("\n".join(header + group))
                            group, group_len = [], header_len
                        group.append(row)
                        group_len += row_len
                    if group:
                        pieces.append("\n".join(header + group))
                continue
            if buf_len + len(seg_text) + 2 <= self.chunk_size or not buf:
                buf.append(seg_text)
                buf_len += len(seg_text) + 2
                if buf_len > self.chunk_size:
                    # 缓冲超长：交给递归切分器细切后直接出货
                    for sub in self.text_splitter.split_text("\n\n".join(buf)):
                        if sub.strip():
                            pieces.append(sub.strip())
                    buf, buf_len = [], 0
            else:
                flush()
                buf.append(seg_text)
                buf_len = len(seg_text)
        flush()
        return [p for p in pieces if p]

    # ------------------------------------------------------------------ #
    # 目录 / 子目录批量加载
    # ------------------------------------------------------------------ #
    def load_directory(
        self,
        dir_path: str | Path,
        recursive: bool = True,
        metadata: dict[str, object] | None = None,
    ) -> list[Document]:
        """
        批量加载目录下的文档。

        :param dir_path: 目录路径
        :param recursive: 是否递归子目录，默认 True
        :param metadata: 可选的额外元数据，透传给每个文件
        :return: 所有文件切分后的片段汇总列表
        """
        path = Path(dir_path)

        # 1. 路径校验：必须存在，且必须是目录
        if not path.exists():
            logger.error("加载失败：目录不存在 -> %s", path)
            return []
        if not path.is_dir():
            logger.error("加载失败：路径不是目录 -> %s", path)
            return []

        logger.info("开始遍历目录：%s | 递归=%s", path, recursive)
        # rglob 本身即递归；非递归时用 glob 只取当前层
        file_iter = path.rglob("*") if recursive else path.glob("*")

        all_docs: list[Document] = []
        scanned = succeeded = failed = 0

        for file_path in sorted(file_iter):
            # 只处理文件，目录本身跳过
            if not file_path.is_file():
                continue
            # 跳过隐藏文件/隐藏目录（.git、.hidden 等）与 Office 临时文件（~$xxx）
            # 按相对路径逐段判断，保证隐藏目录里的文件也被排除
            relative_parts = file_path.relative_to(path).parts
            if any(part.startswith(self.IGNORED_NAME_PREFIXES) for part in relative_parts):
                logger.debug("跳过隐藏/临时文件：%s", file_path)
                continue
            # 跳过不支持的后缀（含无后缀文件）
            if not self.is_supported(file_path):
                logger.debug("跳过不支持的文件：%s", file_path)
                continue

            scanned += 1
            docs = self.load_file(file_path, metadata=metadata)
            if docs:
                succeeded += 1
                all_docs.extend(docs)
            else:
                failed += 1
                logger.warning("文件处理失败已跳过：%s", file_path)

        # 同名文件检测：文件名相同但来源路径不同，内容可能完全不同，必须让使用者知道
        duplicates = self.find_duplicate_names(all_docs)
        if duplicates:
            logger.warning(
                "目录内存在 %d 组同名文件（文件名相同、来源不同，内容可能不一样，请确认是否重复）：%s",
                len(duplicates),
                "；".join(
                    f"{name} -> {', '.join(paths)}" for name, paths in sorted(duplicates.items())
                ),
            )

        logger.info(
            "目录加载完成：%s | 命中文件=%d 成功=%d 失败=%d 片段总数=%d",
            path,
            scanned,
            succeeded,
            failed,
            len(all_docs),
        )
        return all_docs

    # ------------------------------------------------------------------ #
    # 内部方法
    # ------------------------------------------------------------------ #
    def _load_documents(self, path: Path, ext: str) -> list[Document]:
        """按扩展名选择解析策略，返回 Loader 产出的原始文档列表。"""
        # 需要编码回退或结构化解析的后缀，交给专门的分支处理
        if ext in self.CUSTOM_EXTENSIONS:
            return self._load_custom(path, ext)

        # 旧版二进制格式解析库支持有限，提前给出明确告警
        if ext in self.LEGACY_BINARY_FORMATS:
            logger.warning(
                "检测到旧版格式 '%s'：%s 解析能力有限，若失败请先转存为 %s/%s",
                ext,
                path,
                ".docx" if ext == ".doc" else ".xlsx",
                ".xlsx" if ext == ".xls" else ".docx",
            )

        loader_class = self.LOADER_MAP[ext]
        logger.debug("使用 %s 解析文件：%s", loader_class.__name__, path)
        return loader_class(str(path)).load()

    def _load_custom(self, path: Path, ext: str) -> list[Document]:
        """处理需要编码回退或结构化解析的后缀。"""
        # markdown：优先结构化解析，失败时回退为纯文本读取，避免整篇丢失
        if ext == ".md":
            try:
                return UnstructuredMarkdownLoader(str(path)).load()
            except Exception as e:
                logger.warning("Markdown 结构化解析失败（%s），回退为纯文本读取：%s", e, path)
                return self._load_with_encoding_fallback(
                    path, lambda enc: TextLoader(file_path=str(path), encoding=enc)
                )

        # txt：纯文本，逐个编码尝试
        if ext == ".txt":
            return self._load_with_encoding_fallback(
                path, lambda enc: TextLoader(file_path=str(path), encoding=enc)
            )

        # csv：CSVLoader 每行一条 Document（"列名: 值" 拼接），天然适合问答
        if ext == ".csv":
            return self._load_with_encoding_fallback(
                path, lambda enc: CSVLoader(file_path=str(path), encoding=enc)
            )

        # html：取纯文本，用换行分隔避免正文黏成一行
        if ext in {".html", ".htm"}:
            return self._load_with_encoding_fallback(
                path,
                lambda enc: BSHTMLLoader(str(path), open_encoding=enc, get_text_separator="\n"),
            )

        # json：结构未知，统一按「顶层数组 → 每条记录一个 Document，其余 → 整篇一个 Document」
        if ext == ".json":
            return self._load_json(path)

        # pptx：直接用 python-pptx 抽取，每页一条 Document
        if ext == ".pptx":
            return self._load_pptx(path)

        raise ValueError(f"未处理的扩展名: {ext}")

    def _load_pptx(self, path: Path) -> list[Document]:
        """
        用 python-pptx 抽取每页幻灯片的文本（含表格与演讲者备注）。

        不用 UnstructuredPowerPointLoader：它在文本较长时会进入 unstructured 的分块逻辑，
        依赖 NLTK 的 punkt 数据，离线 / 无缓存环境下会抛 LookupError，导致整份 PPT 解析失败。
        """
        presentation = Presentation(str(path))
        documents: list[Document] = []
        for page_no, slide in enumerate(presentation.slides, start=1):
            parts: list[str] = []
            for shape in slide.shapes:
                # 文本框 / 标题 / 占位符（其余形状没有 text_frame，getattr 取不到）
                text_frame = getattr(shape, "text_frame", None)
                if isinstance(text_frame, TextFrame) and text_frame.text.strip():
                    parts.append(text_frame.text.strip())
                # 表格：按行拼接单元格
                table = getattr(shape, "table", None)
                if isinstance(table, Table):
                    for row in table.rows:
                        cells = [cell.text.strip() for cell in row.cells]
                        joined = " | ".join(cell for cell in cells if cell)
                        if joined:
                            parts.append(joined)
            # 演讲者备注往往含关键补充信息，一并纳入
            if slide.has_notes_slide:
                notes_frame = slide.notes_slide.notes_text_frame
                if notes_frame is not None and notes_frame.text.strip():
                    parts.append(f"[备注] {notes_frame.text.strip()}")
            if parts:
                documents.append(Document(page_content="\n".join(parts), metadata={"page": page_no}))
        logger.debug("PPTX 解析完成：%s | 有效页数=%d", path, len(documents))
        return documents

    def _load_json(self, path: Path) -> list[Document]:
        """
        结构化解析 JSON。

        这里刻意不用 JSONLoader：它的 jq_schema 无法同时兼容顶层对象与顶层数组
        —— jq_schema="." 遇到对象会退化成遍历 key（拿到的是字段名而非内容），
        且默认 text_content=True 遇到 dict 会直接抛 ValueError。
        手工处理可保证任意 JSON 结构都完整入库，且不引入额外的 jq 依赖。
        """
        data: object = json.loads(self._read_text(path))
        records: list[object] = data if isinstance(data, list) else [data]
        documents: list[Document] = []
        for record in records:
            content = record if isinstance(record, str) else json.dumps(record, ensure_ascii=False, indent=2)
            if content.strip():
                documents.append(Document(page_content=content))
        logger.debug("JSON 结构化解析完成：%s | 记录数=%d", path, len(documents))
        return documents

    def _read_text(self, path: Path) -> str:
        """按 TEXT_ENCODINGS 依次尝试读取文本文件，全部失败则抛出最后一次异常。"""
        last_error: Exception | None = None
        for encoding in self.TEXT_ENCODINGS:
            try:
                return path.read_text(encoding=encoding)
            except (UnicodeDecodeError, LookupError) as e:
                last_error = e
                logger.debug("以编码 %s 读取失败，尝试下一个：%s", encoding, path)
        logger.error("文本文件解码失败（已尝试 %s）：%s", self.TEXT_ENCODINGS, path)
        raise last_error if last_error else RuntimeError(f"无法解码文本文件: {path}")

    def _load_with_encoding_fallback(
        self, path: Path, build_loader: Callable[[str], BaseLoader]
    ) -> list[Document]:
        """依次用候选编码构造 Loader 并解析，全部失败则抛出最后一次异常。"""
        last_error: Exception | None = None
        for encoding in self.TEXT_ENCODINGS:
            try:
                docs = build_loader(encoding).load()
                logger.debug("以编码 %s 解析成功：%s", encoding, path)
                return docs
            # TextLoader / CSVLoader 会把 UnicodeDecodeError 包装成 RuntimeError，二者都要接住
            except (UnicodeDecodeError, LookupError, RuntimeError) as e:
                last_error = e
                logger.debug("以编码 %s 解析失败，尝试下一个：%s", encoding, path)
        logger.error("文件解码失败（已尝试 %s）：%s", self.TEXT_ENCODINGS, path)
        raise last_error if last_error else RuntimeError(f"无法解码文件: {path}")
