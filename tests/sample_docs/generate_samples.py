"""
生成 document_loader 测试用的样本文件（可重复执行，覆盖式生成）。

设计原则：
- 以 test.txt（企业考勤管理制度）为**唯一内容源**，其余格式都由它派生，保证各格式内容一致、便于断言。
- 不覆盖 test.txt 本身。
- docx 用标准库 zipfile 手工拼装最小 OOXML，避免引入 python-docx 依赖。

运行：
    .venv/bin/python tests/sample_docs/generate_samples.py

依赖：pymupdf（pdf）、openpyxl（xlsx）、python-pptx（pptx）
"""
import csv
import json
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

SAMPLE_DIR = Path(__file__).resolve().parent
NESTED_DIR = SAMPLE_DIR / "nested"
SOURCE_TXT = SAMPLE_DIR / "test.txt"

# 从 test.txt 中提炼的请假类型表，xlsx / csv 共用
LEAVE_TABLE: list[tuple[str, ...]] = [
    ("请假类型", "期限", "工资发放"),
    ("事假", "全年不超过15天，单次不超过5天", "不发放工资"),
    ("病假", "提供医院证明材料，全年累计不超过30天按80%发放", "按累计天数分档发放"),
    ("婚假", "3天，符合晚婚条件合计10天", "基本工资全额发放"),
    ("产假", "98天，难产增加15天，多胞胎每多1个增加15天", "生育津贴，企业补足差额"),
]

# ---------------------------------------------------------------- docx 最小 OOXML
CONTENT_TYPES = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
<Default Extension="xml" ContentType="application/xml"/>
<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
</Types>"""

RELS = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>
</Relationships>"""

DOCUMENT_XML = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
<w:body>
{body}
<w:sectPr/>
</w:body>
</w:document>"""


def read_source_paragraphs() -> list[str]:
    """读取 test.txt 的非空行，作为各格式样本的内容源。"""
    if not SOURCE_TXT.exists():
        raise FileNotFoundError(f"内容源文件不存在：{SOURCE_TXT}")
    lines = SOURCE_TXT.read_text(encoding="utf-8").splitlines()
    return [line.strip() for line in lines if line.strip()]


def write_markdown(paragraphs: list[str]) -> None:
    body = "\n\n".join(paragraphs)
    (SAMPLE_DIR / "test.md").write_text(f"# 企业考勤管理制度\n\n{body}\n", encoding="utf-8")


def write_html(paragraphs: list[str]) -> None:
    items = "\n".join(f"    <p>{escape(p)}</p>" for p in paragraphs)
    html = f"""<!DOCTYPE html>
<html lang="zh-CN">
<head>
    <meta charset="utf-8"/>
    <title>企业考勤管理制度</title>
</head>
<body>
    <h1>企业考勤管理制度</h1>
{items}
</body>
</html>
"""
    (SAMPLE_DIR / "test.html").write_text(html, encoding="utf-8")


def write_json(paragraphs: list[str]) -> None:
    """顶层数组结构：每个条款一条记录。"""
    records: list[dict[str, str]] = []
    current_title = "总则"
    for para in paragraphs:
        # "第一条 xxx" / "第二章 xxx" 视为条款标题
        if para[:2] in {"第一", "第二", "第三", "第四", "第五"} and ("条" in para[:5] or "章" in para[:5]) and len(para) < 30:
            current_title = para
            continue
        records.append({"条款": current_title, "内容": para})
    (SAMPLE_DIR / "test.json").write_text(
        json.dumps(records, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )


def write_csv(rows: list[tuple[str, ...]]) -> None:
    with open(SAMPLE_DIR / "test.csv", "w", encoding="utf-8", newline="") as f:
        csv.writer(f).writerows(rows)


def write_xlsx(rows: list[tuple[str, ...]]) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "请假类型"
    for row in rows:
        ws.append(list(row))
    wb.save(SAMPLE_DIR / "test.xlsx")


def write_pptx() -> None:
    from pptx import Presentation

    prs = Presentation()
    cover = prs.slides.add_slide(prs.slide_layouts[0])
    cover.shapes.title.text = "企业考勤管理制度"
    cover.placeholders[1].text = "请假 / 调休 / 年假 管理规范"

    detail = prs.slides.add_slide(prs.slide_layouts[1])
    detail.shapes.title.text = "请假类型与期限"
    frame = detail.placeholders[1].text_frame
    frame.text = "事假：全年不超过15天，单次不超过5天"
    for line in ("病假：需提供医院证明材料", "婚假：3天，晚婚合计10天", "产假：98天，难产增加15天"):
        frame.add_paragraph().text = line

    prs.save(SAMPLE_DIR / "test.pptx")


def write_pdf(paragraphs: list[str], per_page: int = 10) -> None:
    import pymupdf

    doc = pymupdf.open()
    for i in range(0, len(paragraphs), per_page):
        page = doc.new_page()
        # 内置 CJK 字体 china-s，保证中文可提取（不是图片化文本）
        page.insert_textbox(
            pymupdf.Rect(50, 50, 545, 790),
            "\n".join(paragraphs[i : i + per_page]),
            fontname="china-s",
            fontsize=10,
        )
    doc.save(SAMPLE_DIR / "test.pdf")
    doc.close()


def write_docx(paragraphs: list[str]) -> None:
    body = "\n".join(
        f'<w:p><w:r><w:t xml:space="preserve">{escape(p)}</w:t></w:r></w:p>' for p in paragraphs
    )
    with zipfile.ZipFile(SAMPLE_DIR / "test.docx", "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("[Content_Types].xml", CONTENT_TYPES)
        zf.writestr("_rels/.rels", RELS)
        zf.writestr("word/document.xml", DOCUMENT_XML.format(body=body))


def write_nested(paragraphs: list[str]) -> None:
    NESTED_DIR.mkdir(parents=True, exist_ok=True)
    (NESTED_DIR / "extra.md").write_text(
        "# 附则（子目录样本）\n\n" + "\n\n".join(paragraphs[:3]) + "\n", encoding="utf-8"
    )


def main() -> None:
    paragraphs = read_source_paragraphs()
    write_markdown(paragraphs)
    write_html(paragraphs)
    write_json(paragraphs)
    write_csv(LEAVE_TABLE)
    write_xlsx(LEAVE_TABLE)
    write_pptx()
    write_pdf(paragraphs)
    write_docx(paragraphs)
    write_nested(paragraphs)

    print(f"样本生成完成（内容源：{SOURCE_TXT.name}，段落数 {len(paragraphs)}）：")
    for path in sorted(SAMPLE_DIR.rglob("*")):
        if path.is_file() and path.name != Path(__file__).name:
            print(f"  {path.relative_to(SAMPLE_DIR)}  ({path.stat().st_size} B)")


if __name__ == "__main__":
    main()
